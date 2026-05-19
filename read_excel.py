import openpyxl
wb = openpyxl.load_workbook("/Users/ayralu/Desktop/达能/Y26_Content Hub/功能清单/TQC_Y26_达能_ContentHub_功能清单_用户视角_260420.xlsx", data_only=True)
print("Sheets:", wb.sheetnames)
for sheet in wb.sheetnames:
    ws = wb[sheet]
    print(f"\n=== {sheet} (共 {ws.max_row} 行) ===")
    for row in ws.iter_rows(min_row=1, max_row=ws.max_row, values_only=True):
        if any(v is not None for v in row):
            print(row)
